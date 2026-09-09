"""Spreadsheet extractor -- every sheet becomes one table."""

from __future__ import annotations

from io import BytesIO

import openpyxl

from ar_pipeline.extract.base import ExtractedContent


def extract_excel(data: bytes) -> ExtractedContent:
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        tables: list[list[list[str]]] = []
        blocks: list[str] = []
        for name in sheet_names:
            sheet = workbook[name]
            table: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                if not any(cell is not None for cell in row):
                    continue
                table.append([str(cell) if cell is not None else "" for cell in row])
            tables.append(table)
            body = "\n".join("\t".join(row) for row in table)
            blocks.append(f"### {name}\n{body}")
    finally:
        workbook.close()

    return ExtractedContent(
        text="\n\n".join(blocks),
        tables=tables,
        meta={"sheet_names": sheet_names},
    )
